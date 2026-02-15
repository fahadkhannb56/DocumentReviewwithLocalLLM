# resume_ats_rag.py
# Output: ATS_Report.pdf (and optionally ATS_Report.md)
# Local pipeline: Resume -> embeddings -> Chroma -> retrieval -> Ollama(Mistral) -> PDF

import os
import warnings
from datetime import datetime

# Optional: reduce console noise
warnings.filterwarnings("ignore", message=".*LangChainDeprecationWarning.*")
warnings.filterwarnings("ignore", message=".*CryptographyDeprecationWarning.*")
os.environ.setdefault("TF_CPP_MIN_LOG_LEVEL", "3")

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import Chroma
from langchain_ollama import OllamaLLM

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm


def load_resume(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import docx
        d = docx.Document(path)
        return "\n".join(p.text for p in d.paragraphs)

    if ext == ".txt":
        with open(path, "r", encoding="utf-8") as f:
            return f.read()

    raise ValueError("Unsupported file type. Use PDF, DOCX, or TXT.")


def write_text_pdf(output_path: str, title: str, text: str) -> None:
    """
    Minimal, reliable PDF writer with word-wrapping + page breaks.
    Produces a clean, readable PDF without needing HTML/markdown converters.
    """
    c = canvas.Canvas(output_path, pagesize=A4)
    width, height = A4

    left = 2.0 * cm
    right = 2.0 * cm
    top = 2.0 * cm
    bottom = 2.0 * cm

    max_width = width - left - right
    y = height - top

    # Title
    c.setFont("Helvetica-Bold", 16)
    c.drawString(left, y, title)
    y -= 0.8 * cm

    # Meta line
    c.setFont("Helvetica", 9)
    meta = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Local: Ollama (mistral:7b) | RAG: Chroma"
    c.drawString(left, y, meta)
    y -= 0.7 * cm

    # Body
    c.setFont("Helvetica", 10)
    line_height = 14

    def wrap_line(line: str):
        # simple wrap by words
        words = line.split()
        if not words:
            return [""]
        lines = []
        cur = words[0]
        for w in words[1:]:
            candidate = cur + " " + w
            if c.stringWidth(candidate, "Helvetica", 10) <= max_width:
                cur = candidate
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
        return lines

    for raw_line in text.splitlines():
        # preserve headings visually
        if raw_line.strip().startswith("#"):
            # Markdown heading -> bold line
            heading = raw_line.strip().lstrip("#").strip()
            if y < bottom + 2 * cm:
                c.showPage()
                y = height - top
            c.setFont("Helvetica-Bold", 12)
            for wl in wrap_line(heading):
                c.drawString(left, y, wl)
                y -= line_height
            y -= 4
            c.setFont("Helvetica", 10)
            continue

        # normal line
        wrapped = wrap_line(raw_line.rstrip())
        for wl in wrapped:
            if y < bottom + line_height:
                c.showPage()
                y = height - top
                c.setFont("Helvetica", 10)
            c.drawString(left, y, wl)
            y -= line_height

    c.save()


def main():
    # ---- CONFIG ----
    RESUME_PATH = "resume.pdf"            # adjust if needed
    VECTOR_DIR = "./chroma_resume_db"     # persistent local store
    MODEL_NAME = "mistral:7b"

    EXPORT_MD = False                     # set True if you also want ATS_Report.md
    OUT_MD = "ATS_Report.md"
    OUT_PDF = "ATS_Report.pdf"

    if not os.path.exists(RESUME_PATH):
        raise FileNotFoundError(f"Resume not found: {RESUME_PATH}")

    # ---- LOAD ----
    resume_text = load_resume(RESUME_PATH).strip()
    if not resume_text or len(resume_text) < 300:
        raise RuntimeError("Resume text is empty / not extractable. Use DOCX or text-based PDF.")

    # ---- CHUNK ----
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120)
    docs = splitter.create_documents([resume_text])

    # ---- EMBEDDINGS ----
    embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")

    # ---- VECTOR STORE ----
    vectordb = Chroma.from_documents(
        documents=docs,
        embedding=embeddings,
        persist_directory=VECTOR_DIR
    )

    # ---- LLM ----
    llm = OllamaLLM(model=MODEL_NAME)

    # ---- RETRIEVAL ----
    retriever = vectordb.as_retriever(search_kwargs={"k": 10})
    chunks = retriever.invoke(
        "Full resume including contact, summary, skills, experience, education, certifications, projects"
    )
    context = "\n\n---\n\n".join(d.page_content for d in chunks)

    # ---- PROMPT ----
    prompt = """
You are an ATS resume auditor.
Use ONLY the RESUME CONTEXT below. Do not assume or infer anything not explicitly present.
If information is missing, say: "Not found in resume".

Deliver in clean markdown with headings:

## 1) ATS Score (0–100)
- Provide a weighted rubric totaling 100 (Formatting, Keywords, Impact, Role Fit).
- For each category, justify with evidence from resume context.

## 2) Keyword Gap Analysis
Gaps by category:
- Delivery / PM
- AI / GenAI / ML
- Data Platforms / Analytics
- Governance / Compliance

## 3) ATS Parsing Risks
Headings, bullets, dates consistency, tables/columns, icons/images, hyperlinks, contact placement.

## 4) Impact Rewrites (6 bullets)
Rewrite 6 bullets as: Action + Scope + Outcome + Metric placeholder.
Do not invent metrics; use placeholders like [X%], [AED X], [N].

## 5) Top 10 Improvements (prioritized)
High ROI edits first.
"""

    final_input = f"RESUME CONTEXT:\n{context}\n\nTASK:\n{prompt}"
    output = llm.invoke(final_input)

    # ---- OUTPUTS ----
    if EXPORT_MD:
        with open(OUT_MD, "w", encoding="utf-8") as f:
            f.write(output)
        print(f"ATS report generated: {OUT_MD}")

    write_text_pdf(
        output_path=OUT_PDF,
        title="ATS Resume Report",
        text=output
    )
    print(f"ATS report generated: {OUT_PDF}")


if __name__ == "__main__":
    main()
